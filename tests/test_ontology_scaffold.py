# SPDX-License-Identifier: Apache-2.0

"""The Ontology Manager writes the Entity–Attribute–Relation pattern by
construction, and the readers find it by its triples.

Every operation runs against the REAL vendored core in a fresh workspace, and
the pattern invariant (:func:`backend.ontology_scaffold.check_pattern`) must
hold afterwards. Assertions read the persisted extension with rdflib.
"""
from __future__ import annotations

import shutil

import pytest
from rdflib import OWL, RDF, RDFS, Graph, Literal, Namespace, URIRef
from rdflib.graph import ReadOnlyGraphAggregate

from backend.ontology_kinds import core_graph
from backend.ontology_manager import OntologyFunctions, apply_extension_instructions
from backend.ontology_scaffold import (
    attributes_of, category_members_by_component, category_of, check_pattern,
    general_predicate_of, link_predicate, own_category, own_general_predicate,
    specific_predicates_of,
)
from backend.workspace.storage import WorkspaceStorage

DICI = Namespace("https://digicities.info/ontology#")
EXT = "scaffold.ttl"


@pytest.fixture()
def funcs(tmp_path):
    of = OntologyFunctions(storage=WorkspaceStorage.local(str(tmp_path)), workspace_id="ws")
    assert of.create_new_extension("scaffold")[0]
    assert of.load_extension_and_update(EXT)[0]
    return of


def _view(funcs) -> Graph:
    return ReadOnlyGraphAggregate([funcs.load_extension(EXT), funcs.load_core_ontology()])


def _ok(result):
    ok, msg = result
    assert ok, msg


def _clean(funcs, *names):
    problems = check_pattern(_view(funcs), [DICI[n] for n in names])
    assert problems == [], problems


# -- reading the core -----------------------------------------------------------

def test_core_categories_are_read_from_domain_and_range():
    g = core_graph()
    assert own_general_predicate(g, DICI.Turbine) == DICI.hasTurbineAttribute
    assert own_category(g, DICI.Turbine) == DICI.TurbineAttribute
    # Since v0.6.0 the core's scaffold is built by the Ontology Manager: the
    # category of LiquidFuel follows the class, the old name is a deprecated alias.
    assert own_category(g, DICI.LiquidFuel) == DICI.LiquidFuelAttribute
    assert (DICI.LiquidFuelCarrierAttribute, OWL.equivalentClass,
            DICI.LiquidFuelAttribute) in g
    assert own_category(g, DICI.GasMeter) == DICI.GasMeterAttribute
    # The old short duplicate is an alias now; the full predicate is the general one.
    assert own_general_predicate(g, DICI.ColdCarrier) == DICI.hasColdCarrierAttribute
    # A class with no scaffold of its own inherits its nearest ancestor's.
    bare = Graph()
    bare.add((DICI.OddMeter, RDFS.subClassOf, DICI.GasMeter))
    view = ReadOnlyGraphAggregate([bare, g])
    assert own_category(view, DICI.OddMeter) is None
    assert category_of(view, DICI.OddMeter) == DICI.GasMeterAttribute
    assert general_predicate_of(view, DICI.OddMeter) == DICI.hasGasMeterAttribute


def test_core_satisfies_the_pattern():
    assert check_pattern(core_graph()) == []


# -- operations -----------------------------------------------------------------

def test_add_component_under_a_core_leaf_hangs_under_the_leaf_scaffold(funcs):
    _ok(funcs.add_component(EXT, "Smart Gas Meter", str(DICI.GasMeter)))
    v = _view(funcs)
    assert own_category(v, DICI.GasMeter) == DICI.GasMeterAttribute
    assert (DICI.GasMeterAttribute, RDFS.subClassOf, DICI.MeterAttribute) in v
    assert (DICI.hasGasMeterAttribute, RDFS.subPropertyOf, DICI.hasMeterAttribute) in v
    assert own_category(v, DICI.SmartGasMeter) == DICI.SmartGasMeterAttribute
    assert (DICI.SmartGasMeterAttribute, RDFS.subClassOf, DICI.GasMeterAttribute) in v
    assert set(v.objects(DICI.hasSmartGasMeterAttribute, RDFS.range)) == {
        DICI.SmartGasMeterAttribute}
    _clean(funcs, "GasMeter", "SmartGasMeter")


def test_a_child_of_a_class_without_a_predicate_gets_a_declared_super_property(funcs):
    """The traffic bug: `hasVehicleCountObservationAttribute ⊑ hasObservationAttribute`
    pointed at a property nobody declared, so its edges fell outside hasAttribute."""
    _ok(funcs.add_component(EXT, "Vehicle Count Observation", str(DICI.Observation)))
    v = _view(funcs)
    general = own_general_predicate(v, DICI.VehicleCountObservation)
    parent_general = general_predicate_of(v, DICI.Observation)
    assert (general, RDFS.subPropertyOf, parent_general) in v
    assert (parent_general, RDF.type, OWL.ObjectProperty) in v
    # The core's existing ObservationAttribute is adopted, not duplicated.
    assert own_category(v, DICI.Observation) == DICI.ObservationAttribute
    _clean(funcs, "Observation", "VehicleCountObservation")


def test_link_writes_one_specific_predicate_and_no_attribute_range(funcs):
    _ok(funcs.add_attribute(EXT, "Physical", "Panel Area", qudt_unit="M2"))
    _ok(funcs.add_component(EXT, "Solar Panel", str(DICI.Turbine)))
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))   # idempotent
    v = _view(funcs)
    assert set(v.objects(DICI.hasSolarPanelAttribute, RDFS.range)) == {DICI.SolarPanelAttribute}
    assert specific_predicates_of(v, DICI.SolarPanel) == [DICI.hasSolarPanelPanelAreaAttribute]
    assert (DICI.hasSolarPanelPanelArea, None, None) not in v
    assert attributes_of(v, DICI.SolarPanel) == [DICI.PanelArea]
    assert (DICI.PanelArea, RDFS.subClassOf, DICI.SolarPanelAttribute) in v
    assert funcs.get_component_attributes(EXT, str(DICI.SolarPanel)) == [
        {"class": str(DICI.PanelArea), "label": "Panel Area"}]
    _clean(funcs, "SolarPanel")


def test_an_attribute_of_an_ancestor_is_linked_through_the_ancestors_predicate(funcs):
    _ok(funcs.add_attribute(EXT, "Physical", "Hub Height", qudt_unit="M"))
    _ok(funcs.link_attribute(EXT, str(DICI.Turbine), str(DICI.HubHeight)))
    _ok(funcs.add_component(EXT, "Wind Turbine", str(DICI.Turbine)))
    v = _view(funcs)
    assert link_predicate(v, DICI.WindTurbine, DICI.HubHeight) == DICI.hasTurbineHubHeightAttribute
    assert link_predicate(v, DICI.WindTurbine, DICI.Efficiency) is None


def test_move_carries_the_category_and_general_predicate(funcs):
    _ok(funcs.add_attribute(EXT, "Physical", "Panel Area", qudt_unit="M2"))
    _ok(funcs.add_component(EXT, "Solar Panel", str(DICI.Turbine)))
    _ok(funcs.add_component(EXT, "Bifacial Panel", str(DICI.SolarPanel)))
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    _ok(funcs.change_component_parent(EXT, str(DICI.SolarPanel), str(DICI.Storage)))
    v = _view(funcs)
    assert set(v.objects(DICI.SolarPanelAttribute, RDFS.subClassOf)) == {DICI.StorageAttribute}
    assert set(v.objects(DICI.hasSolarPanelAttribute, RDFS.subPropertyOf)) == {
        DICI.hasStorageAttribute}
    # The attribute and the subclass's scaffold follow untouched.
    assert (DICI.PanelArea, RDFS.subClassOf, DICI.SolarPanelAttribute) in v
    assert (DICI.BifacialPanelAttribute, RDFS.subClassOf, DICI.SolarPanelAttribute) in v
    _clean(funcs, "SolarPanel", "BifacialPanel")


def test_move_back_up_rehangs_the_category_directly(funcs):
    """Moved from a class to that class's ancestor, the category must hang
    directly under the ancestor's category, not stay under the old parent's."""
    _ok(funcs.add_component(EXT, "Rotor", str(DICI.Turbine)))
    _ok(funcs.add_component(EXT, "Blade", str(DICI.Rotor)))
    _ok(funcs.change_component_parent(EXT, str(DICI.Blade), str(DICI.Turbine)))
    v = _view(funcs)
    assert set(v.objects(DICI.BladeAttribute, RDFS.subClassOf)) == {DICI.TurbineAttribute}
    _clean(funcs, "Rotor", "Blade")


def test_rename_follows_the_scaffold_by_its_triples(funcs):
    _ok(funcs.add_attribute(EXT, "Physical", "Panel Area", qudt_unit="M2"))
    _ok(funcs.add_component(EXT, "Solar Panel", str(DICI.Turbine)))
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    _ok(funcs.rename_component(EXT, str(DICI.SolarPanel), "PV Panel"))
    v = _view(funcs)
    assert own_category(v, DICI.PVPanel) == DICI.PVPanelAttribute
    assert own_general_predicate(v, DICI.PVPanel) == DICI.hasPVPanelAttribute
    assert specific_predicates_of(v, DICI.PVPanel) == [DICI.hasPVPanelPanelAreaAttribute]
    assert (DICI.PVPanelAttribute, RDFS.label, Literal("PVPanel Attribute")) in v
    for gone in (DICI.SolarPanel, DICI.SolarPanelAttribute, DICI.hasSolarPanelAttribute,
                 DICI.hasSolarPanelPanelAreaAttribute):
        assert (gone, None, None) not in v and (None, None, gone) not in v
    _clean(funcs, "PVPanel")


def test_unlink_removes_exactly_what_link_wrote(funcs):
    _ok(funcs.add_attribute(EXT, "Physical", "Panel Area", qudt_unit="M2"))
    _ok(funcs.add_component(EXT, "Solar Panel", str(DICI.Turbine)))
    before = set(funcs.load_extension(EXT))
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    _ok(funcs.remove_attribute_link(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    assert set(funcs.load_extension(EXT)) == before
    ok, msg = funcs.remove_attribute_link(EXT, str(DICI.SolarPanel), str(DICI.PanelArea))
    assert not ok and "not linked" in msg


def test_round_trip_returns_the_extension_to_where_it_started(funcs):
    """Parents whose core chain already states its scaffold (Storage, EnergyStorage):
    nothing is minted for an ancestor, so removing the component leaves nothing."""
    _ok(funcs.add_attribute(EXT, "Physical", "Panel Area", qudt_unit="M2"))
    start = set(funcs.load_extension(EXT))
    _ok(funcs.add_component(EXT, "Solar Panel", str(DICI.Storage)))
    _clean(funcs, "SolarPanel")
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    _clean(funcs, "SolarPanel")
    _ok(funcs.change_component_parent(EXT, str(DICI.SolarPanel), str(DICI.EnergyStorage)))
    _clean(funcs, "SolarPanel")
    _ok(funcs.rename_component(EXT, str(DICI.SolarPanel), "PV Panel"))
    _clean(funcs, "PVPanel")
    _ok(funcs.remove_attribute_link(EXT, str(DICI.PVPanel), str(DICI.PanelArea)))
    _clean(funcs, "PVPanel")
    _ok(funcs.remove_component(EXT, str(DICI.PVPanel)))
    assert set(funcs.load_extension(EXT)) == start


def test_remove_attribute_takes_its_link_predicates(funcs):
    _ok(funcs.add_attribute(EXT, "Physical", "Panel Area", qudt_unit="M2"))
    _ok(funcs.add_component(EXT, "Solar Panel", str(DICI.Turbine)))
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    _ok(funcs.remove_attribute(EXT, str(DICI.PanelArea)))
    v = _view(funcs)
    assert (DICI.hasSolarPanelPanelAreaAttribute, None, None) not in v
    assert (DICI.hasSolarPanelAttribute, RDF.type, OWL.ObjectProperty) in v
    _clean(funcs, "SolarPanel")


def test_category_members_are_read_without_building_names(funcs):
    _ok(funcs.add_attribute(EXT, "Physical", "Panel Area", qudt_unit="M2"))
    _ok(funcs.add_component(EXT, "Solar Panel", str(DICI.Turbine)))
    _ok(funcs.add_component(EXT, "Bifacial Panel", str(DICI.SolarPanel)))
    _ok(funcs.link_attribute(EXT, str(DICI.SolarPanel), str(DICI.PanelArea)))
    members = category_members_by_component(funcs.merge_ontologies(EXT))
    assert members[DICI.SolarPanel] == [DICI.PanelArea]          # no sub-category
    assert members[DICI.BifacialPanel] == []


def test_an_older_extension_is_repaired_when_its_component_is_touched(funcs):
    """The pattern an older Ontology Manager wrote: the general predicate with the
    linked attribute as its range, no category range. Linking again through the
    tooling leaves the category as the only range."""
    g = funcs.load_extension(EXT)
    for t in ((DICI.Pump, RDF.type, OWL.Class), (DICI.Pump, RDFS.subClassOf, DICI.Turbine),
              (DICI.PumpAttribute, RDF.type, OWL.Class),
              (DICI.PumpAttribute, RDFS.subClassOf, DICI.TurbineAttribute),
              (DICI.Head, RDF.type, OWL.Class), (DICI.Head, RDFS.subClassOf, DICI.PumpAttribute),
              (DICI.Head, RDFS.subClassOf, DICI.PhysicalAttribute),
              (DICI.hasPumpAttribute, RDF.type, OWL.ObjectProperty),
              (DICI.hasPumpAttribute, RDFS.subPropertyOf, DICI.hasTurbineAttribute),
              (DICI.hasPumpAttribute, RDFS.domain, DICI.Pump),
              (DICI.hasPumpAttribute, RDFS.range, DICI.Head),
              (DICI.hasPumpHead, RDF.type, OWL.ObjectProperty),
              (DICI.hasPumpHead, RDFS.subPropertyOf, DICI.hasPumpAttribute),
              (DICI.hasPumpHead, RDFS.domain, DICI.Pump),
              (DICI.hasPumpHead, RDFS.range, DICI.Head)):
        g.add(t)
    assert funcs.save_extension(EXT, g)
    funcs.update_temp_and_export(EXT)
    v = _view(funcs)
    assert own_category(v, DICI.Pump) is None          # the attribute is not a category
    assert check_pattern(v, [DICI.Pump])               # broken before
    _ok(funcs.link_attribute(EXT, str(DICI.Pump), str(DICI.Head)))
    v = _view(funcs)
    assert set(v.objects(DICI.hasPumpAttribute, RDFS.range)) == {DICI.PumpAttribute}


# -- the instruction executor ---------------------------------------------------

def _instructions(ops, extension=EXT):
    return {"version": 1, "workspace": "ws", "extension": extension, "instructions": ops}


def test_annotate_is_idempotent_keeps_the_language_and_takes_property_iris(funcs, tmp_path):
    storage = WorkspaceStorage.local(str(tmp_path))
    ops = [
        {"op": "add_component", "name": "SolarPanel", "parent": "Turbine",
         "annotations": {"comment": "First comment."}},
        {"op": "annotate", "name": "hasSolarPanelAttribute",
         "annotations": {"label": "has solar panel attribute", "label_lang": "en",
                         "comment": "Links a solar panel to its attributes."}},
    ]
    for _ in range(2):
        report = apply_extension_instructions(_instructions(ops), storage=storage,
                                              workspace_id="ws")
        assert report["ok"], report["results"]
    ops[0]["op"] = "annotate"
    ops[0]["annotations"] = {"comment": "Second comment.", "label": "Solar panel"}
    report = apply_extension_instructions(_instructions(ops[:1]), storage=storage,
                                          workspace_id="ws")
    assert report["ok"], report["results"]
    g = funcs.load_extension(EXT)
    assert list(g.objects(DICI.SolarPanel, RDFS.comment)) == [Literal("Second comment.",
                                                                      lang="en")]
    assert list(g.objects(DICI.SolarPanel, RDFS.label)) == [Literal("Solar panel")]
    assert list(g.objects(DICI.hasSolarPanelAttribute, RDFS.label)) == [
        Literal("has solar panel attribute", lang="en")]
    assert len(list(g.objects(DICI.hasSolarPanelAttribute, RDFS.comment))) == 1
    # A replaced label keeps the language of the one it replaces.
    report = apply_extension_instructions(_instructions([
        {"op": "annotate", "name": "hasSolarPanelAttribute",
         "annotations": {"label": "has PV panel attribute"}}]), storage=storage,
        workspace_id="ws")
    assert report["ok"], report["results"]
    assert list(funcs.load_extension(EXT).objects(DICI.hasSolarPanelAttribute, RDFS.label)) == [
        Literal("has PV panel attribute", lang="en")]


def test_instructions_can_target_the_core(tmp_path, monkeypatch):
    onto = tmp_path / "onto"
    shutil.copytree(OntologyFunctions._resolve_global_ontology_dir(), onto)
    monkeypatch.setenv("ONTOLOGY_DIR", str(onto))
    ws = tmp_path / "ws"
    ws.mkdir()
    storage = WorkspaceStorage.local(str(ws))
    report = apply_extension_instructions(_instructions([
        {"op": "add_component", "name": "SmartGasMeter", "parent": "GasMeter"},
        {"op": "annotate", "name": "Turbine", "annotations": {"label": "Turbine (rotary)"}},
    ], extension=OntologyFunctions.CORE_TARGET), storage=storage, workspace_id="ws")
    assert report["ok"], report["results"]
    core = Graph().parse(onto / "dici_onto_core.ttl", format="turtle")
    assert (DICI.SmartGasMeter, RDFS.subClassOf, DICI.GasMeter) in core
    assert own_category(core, DICI.SmartGasMeter) == DICI.SmartGasMeterAttribute
    # The core's labels carry no language tag; the replacement keeps none.
    assert list(core.objects(DICI.Turbine, RDFS.label)) == [Literal("Turbine (rotary)")]
    assert check_pattern(core, [DICI.GasMeter, DICI.SmartGasMeter]) == []


def test_core_build_ops_alias_definitions_and_units(tmp_path, monkeypatch):
    """The ops the core build needs: a deprecated alias (class and property),
    dici_onto:definition, a component's default unit, a plain label."""
    onto = tmp_path / "onto"
    shutil.copytree(OntologyFunctions._resolve_global_ontology_dir(), onto)
    monkeypatch.setenv("ONTOLOGY_DIR", str(onto))
    ws = tmp_path / "ws"
    ws.mkdir()
    ops = [
        {"op": "add_component", "name": "SmartGasMeter", "parent": "GasMeter",
         "annotations": {"label": "Smart gas meter", "label_lang": None,
                         "dici_definition": "A gas meter that reports by itself.",
                         "default_unit": "M3"}},
        {"op": "add_deprecated_alias", "name": "SmartGasMeterCategory",
         "replaced_by": "SmartGasMeterAttribute",
         "annotations": {"label": "Smart Gas Meter Category", "comment": "Old name."}},
        {"op": "add_deprecated_alias", "name": "hasSmartMeterAttribute",
         "replaced_by": "hasSmartGasMeterAttribute",
         "annotations": {"label": "has smart meter attribute", "comment": "Old name."}},
    ]
    for _ in range(2):                                     # a rerun changes nothing
        report = apply_extension_instructions(
            _instructions(ops, extension=OntologyFunctions.CORE_TARGET),
            storage=WorkspaceStorage.local(str(ws)), workspace_id="ws")
        assert all(r["status"] in ("applied", "skipped") for r in report["results"]), \
            report["results"]
    core = Graph().parse(onto / "dici_onto_core.ttl", format="turtle")
    assert list(core.objects(DICI.SmartGasMeter, RDFS.label)) == [Literal("Smart gas meter")]
    assert list(core.objects(DICI.SmartGasMeter, DICI.definition)) == [
        Literal("A gas meter that reports by itself.", lang="en")]
    assert list(core.objects(DICI.SmartGasMeter, DICI.hasDefaultUnit)) == [
        URIRef("http://qudt.org/vocab/unit/M3")]
    assert (DICI.SmartGasMeterCategory, OWL.equivalentClass, DICI.SmartGasMeterAttribute) in core
    assert (DICI.SmartGasMeterCategory, OWL.deprecated, Literal(True)) in core
    assert (DICI.hasSmartMeterAttribute, OWL.equivalentProperty,
            DICI.hasSmartGasMeterAttribute) in core
    # Data still using the old predicate stays under hasAttribute.
    assert (DICI.hasSmartMeterAttribute, RDFS.subPropertyOf,
            DICI.hasSmartGasMeterAttribute) in core
    assert check_pattern(core, [DICI.SmartGasMeter]) == []


# -- the replica converter ------------------------------------------------------

def _hub_height_workbook(path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet = wb.create_sheet("WindTurbine")
    for c, (name, atype) in enumerate([("id", None), ("HubHeight", "SimpleValue")], start=1):
        sheet.cell(row=1, column=c, value=name)
        sheet.cell(row=2, column=c, value=atype)
    sheet.cell(row=7, column=1, value="T1")
    sheet.cell(row=7, column=2, value=120)
    wb.save(path)
    return path


def test_converter_links_values_by_the_declared_predicate(funcs, tmp_path):
    from backend.replica_builder.utils.create_class_and_attribute_graph import (
        process_excel_to_ttl)
    _ok(funcs.add_attribute(EXT, "SimpleValue", "Hub Height"))
    _ok(funcs.add_component(EXT, "Wind Turbine", str(DICI.Turbine)))
    xlsx = _hub_height_workbook(tmp_path / "wb.xlsx")
    ttl = tmp_path / "wb.ttl"
    # Not linked yet: the converter names the column instead of guessing a predicate.
    with pytest.raises(ValueError, match="WindTurbine.HubHeight"):
        process_excel_to_ttl("https://x.org/p", str(xlsx), str(ttl),
                             ontology=funcs.load_extension(EXT))
    # Linked to the ancestor: the ancestor's declared predicate carries the value.
    _ok(funcs.link_attribute(EXT, str(DICI.Turbine), str(DICI.HubHeight)))
    process_excel_to_ttl("https://x.org/p", str(xlsx), str(ttl),
                         ontology=funcs.load_extension(EXT))
    data = Graph().parse(ttl, format="turtle")
    t1 = URIRef("https://x.org/p/WindTurbine/T1")
    assert {p for p, _ in data.predicate_objects(t1)} >= {DICI.hasTurbineHubHeightAttribute}
    assert (t1, DICI.hasWindTurbineHubHeightAttribute, None) not in data
