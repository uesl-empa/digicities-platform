# SPDX-License-Identifier: Apache-2.0

"""Test support: declare a replica workbook's classes and attribute links
through the Ontology Manager, the way a workspace gets them in real use.

The converter links each value by the predicate the workspace ontology
declares for its class and attribute, so every converter test needs the
schema its workbook uses. This builds it with the real instruction executor
(``add_component`` / ``add_attribute`` / ``link_attribute``), never by writing
predicate triples by hand.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from rdflib import Graph

from backend.ontology_kinds import AttributeKind
from backend.ontology_manager import apply_extension_instructions
from backend.replica_builder.attribute_rules import (
    TIME_SERIES_COLUMNS, WorkbookColumn, parse_column_type,
)
from backend.workspace.storage import WorkspaceStorage

# Sheets that hold no class: the citation sheet and the dropdown lists.
_NOT_CLASSES = {"Reference", "Data Validation"}
# Columns the converter writes without a declared attribute link.
_UNLINKED = {WorkbookColumn.CLASS_OBJECT, AttributeKind.IDENTIFIER, AttributeKind.ANNOTATION}
# The Ontology Manager's form labels for the kinds it names with spaces.
_OM_LABEL = {AttributeKind.SIMPLE_COST: "Simple Cost",
             AttributeKind.UNIT_BASED_COST: "Unit-Based Cost"}


def _attribute_op(name: str, col_type, unit, unit_y) -> dict:
    parent = None
    if col_type in TIME_SERIES_COLUMNS or col_type is AttributeKind.DYNAMIC:
        kind, parent = (AttributeKind.PHYSICAL if unit else AttributeKind.SIMPLE_VALUE), \
            "DynamicAttribute"
    elif col_type is AttributeKind.RESOURCE:
        kind, parent = AttributeKind.SIMPLE_VALUE, "DataPathAttribute"
    elif col_type is AttributeKind.IDENTIFIER:
        kind, parent = AttributeKind.SIMPLE_VALUE, "IdentifierAttribute"
    else:
        kind = col_type
    op = {"op": "add_attribute", "name": name, "type": _OM_LABEL.get(kind, kind.value)}
    if parent:
        op["parent"] = parent
    if kind in (AttributeKind.PHYSICAL, AttributeKind.GEOSPATIAL, AttributeKind.UNIT_BASED_COST):
        op["qudt_unit"] = unit or "UNITLESS"
    elif kind is AttributeKind.CURVE:
        op["qudt_unit"] = unit or "UNITLESS"
        op["y_qudt_unit"] = unit_y or "UNITLESS"
    elif kind is AttributeKind.CUSTOM_PHYSICAL_RATIO:
        op["x_unit"] = unit or "UNITLESS"
        op["y_qudt_unit"] = unit_y or "UNITLESS"
    elif kind is AttributeKind.EVENT:
        op["temporal_precision"] = "DateTime"
    return op


def _data_rows(ws) -> list:
    """The instance rows below the header block (six header rows, seven when
    the sheet carries the LinkedClassObjectType row)."""
    rows = list(ws.iter_rows(values_only=True))
    seventh = rows[6] if len(rows) > 6 else ()
    first = 7 if "LinkedClassObjectType" in [str(v).strip() for v in seventh if v] else 6
    return [r for r in rows[first:] if r and r[0] not in (None, "")]


def declare_workbook(xlsx, workspace_dir) -> Graph:
    """Declare every class and attribute link ``xlsx`` uses in a fresh
    extension under ``workspace_dir``, and each categorical cell value as a
    value of its attribute labelled with the cell text (as the template import
    does); return that extension graph."""
    from backend.ontology_manager.naming import class_name

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ops, seen = [], set()
    try:
        for ws in wb.worksheets:
            if ws.title in _NOT_CLASSES:
                continue
            header = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
            names, types, units, units_y = (list(r) for r in header + [()] * (4 - len(header)))
            data = _data_rows(ws)
            ops.append({"op": "add_component", "name": ws.title})
            for i, name in enumerate(names):
                if i == 0 or not name:
                    continue
                col_type = parse_column_type(types[i] if i < len(types) else None,
                                             f"{ws.title}.{name}")
                if col_type is None or col_type in _UNLINKED:
                    continue
                if name not in seen:
                    seen.add(name)
                    ops.append(_attribute_op(name, col_type,
                                             units[i] if i < len(units) else None,
                                             units_y[i] if i < len(units_y) else None))
                ops.append({"op": "link_attribute", "component": ws.title, "attribute": name})
                if col_type is AttributeKind.CATEGORICAL:
                    for value in sorted({str(r[i]).strip() for r in data
                                         if i < len(r) and r[i] not in (None, "")}):
                        if ("value", value) not in seen:
                            seen.add(("value", value))
                            ops.append({"op": "add_named_individual",
                                        "name": class_name(value), "attribute": name,
                                        "annotations": {"label": value}})
    finally:
        wb.close()
    root = Path(workspace_dir)
    report = apply_extension_instructions(
        {"extension": "workbook.ttl", "instructions": ops},
        storage=WorkspaceStorage.local(str(root)), workspace_id="ws")
    errors = [r for r in report["results"] if r["status"] == "error"]
    assert not errors, errors
    return Graph().parse(root / "ontology" / "extensions" / "workbook.ttl", format="turtle")
